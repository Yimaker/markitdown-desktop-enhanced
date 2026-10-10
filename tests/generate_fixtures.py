"""Build synthetic PDF/Office fixtures with no human or machine identity."""
from pathlib import Path
import os

from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
import xlwt

MARKER = 'SYNTHETICDOCUMENTCONTENT'

def plain_pdf(path):
    c = canvas.Canvas(str(path))
    c.setTitle('Synthetic document')
    c.setAuthor('')
    c.setSubject('Anonymous generated test fixture')
    for number in range(2):
        c.setFont('Helvetica', 14)
        c.drawString(64, 750, MARKER)
        c.drawString(64, 720, 'Example value: 42; generated page ' + str(number + 1))
        c.showPage()
    c.save()

def math_pdf(path):
    # Use an installed system font for tests; the font file is never redistributed.
    font_path = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / 'cambria.ttc'
    if not font_path.is_file():
        raise RuntimeError('Math fixture generation needs the Cambria Math system font on Windows.')
    pdfmetrics.registerFont(TTFont('CambriaMath', str(font_path), subfontIndex=1))
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    c = canvas.Canvas(str(path))
    c.setTitle('Synthetic mathematical example')
    c.setAuthor('')
    c.setSubject('Generated symbols and tables; no personal report')

    def text(x, y, value, font='Helvetica', size=12):
        c.setFont(font, size)
        c.drawString(x, y, value)

    def equation(x, y, chunks):
        for value, size, offset in chunks:
            text(x, y + offset, value, 'CambriaMath', size)
            x += pdfmetrics.stringWidth(value, 'CambriaMath', size)
        return x

    def table(y, header):
        x = 64
        width = 380
        height = 30
        c.setLineWidth(0.5)
        for column in range(3):
            c.line(x + column * width / 2, y - 2 * height, x + column * width / 2, y + height)
        for row in range(4):
            c.line(x, y + height - row * height, x + width, y + height - row * height)
        for row in range(3):
            for column in range(2):
                value = header[column] if row == 0 else ('Generated row ' + str(row) if column == 0 else str(row * 10))
                text(x + column * width / 2 + 10, y - row * height + 10, value)

    def page(number):
        text(64, 820, 'SYNTHETIC_PAGE_HEADER')
        text(64, 780, 'SYNTHETIC_MATH_EXAMPLE', size=18)
        text(64, 20, '第 ' + str(number) + ' 页', 'STSong-Light', 9)

    page(1)
    text(64, 750, '问题 1 合成指数示例', 'STSong-Light')
    equation(150, 705, [('u',12,0),('n',8,-3),('-',12,0),('l',12,0),('n',8,-3),('=2',12,0),('-n',8,5)])
    equation(150, 665, [('x',12,0),('2',8,5),('=2',12,0)])
    text(64, 565, '问题 2 合成下标示例', 'STSong-Light')
    equation(150, 525, [('a',12,0),('m',8,-3),('-',12,0),('a',12,0),('n',8,-3)])
    equation(150, 485, [('2',12,0),('-(n+1)',8,5)])
    text(64, 405, '问题 3 合成复合下标', 'STSong-Light')
    equation(150, 365, [('l',12,0),('n+1',8,-3)])
    table(265, ['Generated item', 'Example value'])
    c.showPage()
    page(2)
    text(64, 750, '问题 4 合成分数示例', 'STSong-Light')
    end = equation(150, 698, [('u',12,0),('n',8,-3),('-',12,0),('l',12,0),('n',8,-3)])
    c.setLineWidth(0.5)
    c.line(148, 687, end + 2, 687)
    text((150 + end) / 2 - 3, 671, '2', 'CambriaMath', 12)
    text(64, 565, '问题 5 合成表格示例', 'STSong-Light')
    table(505, ['Generated index', 'Synthetic amount'])
    text(64, 335, '问题 6 合成文本示例', 'STSong-Light')
    text(64, 300, 'All values in this document are generated examples.')
    c.showPage()
    c.save()

def generate(root):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    pdf = root / 'Synthetic example.pdf'
    plain_pdf(pdf)
    doc = Document()
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = 'Synthetic document'
    doc.add_heading('Synthetic document', 0)
    doc.add_paragraph(MARKER)
    doc.save(root / 'Synthetic example.docx')
    presentation = Presentation()
    presentation.core_properties.author = ''
    presentation.core_properties.last_modified_by = ''
    presentation.core_properties.title = 'Synthetic slides'
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = 'Synthetic slides'
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1)).text_frame.text = MARKER
    presentation.save(root / 'Synthetic example.pptx')
    book = Workbook()
    book.properties.creator = 'Synthetic test generator'
    book.properties.lastModifiedBy = ''
    sheet = book.active
    sheet.title = 'Generated'
    sheet.append(['Synthetic label', 'Example value'])
    sheet.append([MARKER, 42])
    book.save(root / 'Synthetic example.xlsx')
    legacy = xlwt.Workbook()
    legacy_sheet = legacy.add_sheet('Generated')
    legacy_sheet.write(0, 0, 'Synthetic label')
    legacy_sheet.write(0, 1, 'Example value')
    legacy_sheet.write(1, 0, MARKER)
    legacy_sheet.write(1, 1, 42)
    legacy.save(str(root / 'Synthetic example.xls'))
    mathematical = root / 'Synthetic math example.pdf'
    math_pdf(mathematical)
    return {'official': [{'path': str(root / ('Synthetic example.' + extension)),
                         'must_include': [MARKER], 'must_not_include': []}
                        for extension in ('pdf', 'docx', 'pptx', 'xlsx', 'xls')],
            'reference_pdf': str(mathematical)}
